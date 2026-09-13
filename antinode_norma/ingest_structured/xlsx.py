from pathlib import Path
from typing import List, Union, Dict, Any, Optional
import openpyxl
from antinode_norma.core.types import TestCase
from antinode_norma.core.normalize import normalize_header, split_multi


class XLSXIngester:
    """Ingests XLSX workbook worksheets into standardized TestCase instances."""

    def __init__(self, sheet_name: Optional[str] = None):
        self.sheet_name = sheet_name

    def ingest(self, path: Union[str, Path]) -> List[TestCase]:
        xlsx_path = Path(path)
        if not xlsx_path.exists():
            raise FileNotFoundError(f"XLSX file not found: {xlsx_path}")

        wb = openpyxl.load_workbook(filename=xlsx_path, data_only=True)
        if self.sheet_name and self.sheet_name in wb.sheetnames:
            ws = wb[self.sheet_name]
        else:
            ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        header_row = [str(cell) if cell is not None else "" for cell in rows[0]]
        header_map: Dict[int, str] = {
            col_idx: normalize_header(raw)
            for col_idx, raw in enumerate(header_row)
            if raw
        }

        test_cases: List[TestCase] = []

        for idx, row in enumerate(rows[1:], 1):
            if not any(cell is not None for cell in row):
                continue

            mapped_row: Dict[str, Any] = {}
            metadata: Dict[str, Any] = {}

            for col_idx, cell_val in enumerate(row):
                if cell_val is None or col_idx not in header_map:
                    continue
                norm_key = header_map[col_idx]
                val = str(cell_val).strip()

                if norm_key in {
                    "id",
                    "title",
                    "role",
                    "action",
                    "benefit",
                    "acceptance_criteria",
                    "tags",
                }:
                    mapped_row[norm_key] = val
                else:
                    metadata[norm_key] = val

            case_id = str(mapped_row.get("id") or f"TC-{idx:03d}").strip()
            title = str(mapped_row.get("title") or f"Test Case {idx}").strip()
            role = str(mapped_row.get("role") or "user").strip()
            action = str(mapped_row.get("action") or "").strip()
            benefit = str(mapped_row.get("benefit") or "").strip()

            raw_criteria = mapped_row.get("acceptance_criteria")
            acceptance_criteria = split_multi(raw_criteria)

            raw_tags = mapped_row.get("tags")
            tags = split_multi(raw_tags)

            test_cases.append(
                TestCase(
                    id=case_id,
                    title=title,
                    role=role,
                    action=action,
                    benefit=benefit,
                    acceptance_criteria=acceptance_criteria,
                    tags=tags,
                    metadata=metadata,
                )
            )

        return test_cases
